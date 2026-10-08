import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QPoint, QPointF, QRect, QSize, Qt
from PySide6.QtGui import QWheelEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QFileDialog

from app.geometry import bubble_position
from app.library import LibraryStore
from app.models import ExtensionAction
from app.questions import SAMPLE_QUESTIONS, question_to_dict

ROOT = Path(__file__).resolve().parents[1]


def hover_info(application, text):
    question = replace(SAMPLE_QUESTIONS[0], extension=ExtensionAction("info", text))
    application.controller.load_questions((question,))
    application.overlay.show()
    widget = application.overlay.widget
    QTest.mouseMove(widget, widget.rect().center())
    QTest.mouseMove(widget, widget.corner_centers()[0])
    QApplication.processEvents()
    return widget, application.action_context._info_dialog


def test_info_bubble_shows_only_while_hovering_question_mark(application):
    widget, bubble = hover_info(application, "<b>纯文本解析</b>")
    assert bubble.isVisible()
    assert bubble.text.toPlainText() == "<b>纯文本解析</b>"
    assert bubble.windowType() == Qt.WindowType.ToolTip
    assert not application.session.selected
    QTest.mouseMove(widget, widget.rect().center())
    assert not bubble.isVisible()
    QTest.mouseMove(widget, widget.corner_centers()[0])
    assert bubble.isVisible()
    application.overlay.hide()
    assert not bubble.isVisible()


def test_info_wheel_scrolls_in_place_and_theme_changes_immediately(application):
    widget, bubble = hover_info(application, "\n".join(f"解析第 {i} 行" for i in range(100)))
    assert bubble.text.verticalScrollBar().maximum() > 0
    point = widget.corner_centers()[0]
    event = QWheelEvent(
        QPointF(point),
        QPointF(widget.mapToGlobal(point)),
        QPoint(),
        QPoint(0, -120),
        Qt.MouseButton.NoButton,
        Qt.KeyboardModifier.NoModifier,
        Qt.ScrollPhase.NoScrollPhase,
        False,
    )
    QApplication.sendEvent(widget, event)
    assert bubble.text.verticalScrollBar().value() > 0
    assert bubble.isVisible()
    from PySide6.QtGui import QColor

    application.preferences.set_color("background", QColor("#FF203040"))
    application.apply_settings()
    assert bubble.theme == application.preferences.colors
    pixel = bubble.grab().toImage().pixelColor(bubble.width() // 2, 12)
    assert pixel == QColor("#FF203040")


@pytest.mark.parametrize("background", ["#00121A28", "#20203040"])
def test_info_background_preserves_the_radial_transparency(application, background):
    from PySide6.QtGui import QColor

    application.preferences.set_color("background", QColor(background))
    application.apply_settings()
    _widget, bubble = hover_info(application, "补充解析")
    pixel = bubble.grab().toImage().pixelColor(bubble.width() // 2, 12)
    assert pixel.alpha() == QColor(background).alpha()


def test_info_size_follows_content_and_long_text_remains_scrollable(application):
    _widget, bubble = hover_info(application, "一句补充说明")
    assert bubble.height() <= 96
    assert bubble.text.verticalScrollBar().maximum() == 0
    _widget, bubble = hover_info(application, "\n".join(f"说明 {i}" for i in range(100)))
    assert bubble.height() <= application.overlay.height()
    scrollbar = bubble.text.verticalScrollBar()
    assert scrollbar.maximum() > 0
    scrollbar.setValue(scrollbar.maximum())
    assert scrollbar.value() == scrollbar.maximum()


def test_url_hover_does_not_open_it_but_click_does(application, monkeypatch):
    opened = []
    monkeypatch.setattr(application.action_context, "open_url", opened.append)
    question = replace(SAMPLE_QUESTIONS[0], extension=ExtensionAction("url", "https://example.com"))
    application.controller.load_questions((question,))
    application.overlay.show()
    widget = application.overlay.widget
    QTest.mouseMove(widget, widget.corner_centers()[0])
    assert opened == []
    assert application.action_context._info_dialog is None
    QTest.mouseClick(widget, Qt.MouseButton.LeftButton, pos=widget.corner_centers()[0])
    assert opened == ["https://example.com"]


@pytest.mark.parametrize(
    "anchor",
    [
        QRect(-1910, 10, 240, 240),
        QRect(-260, 10, 240, 240),
        QRect(-1910, 800, 240, 240),
        QRect(-260, 800, 240, 240),
        QRect(-1080, 420, 240, 240),
    ],
)
def test_bubble_placement_stays_on_monitor_and_clear_of_options(anchor):
    available = QRect(-1920, 0, 1920, 1080)
    size = QSize(360, 390)
    position, direction = bubble_position(anchor, available, size)
    bubble = QRect(position, size)
    assert direction in {"left", "right", "top", "bottom"}
    assert available.contains(bubble)
    assert not bubble.intersects(anchor)


def test_matching_ids_in_different_files_keep_separate_progress(application, tmp_path):
    sources = []
    question = replace(SAMPLE_QUESTIONS[0], id="1")
    for folder in ("english", "math"):
        source = tmp_path / folder / "bank.json"
        source.parent.mkdir()
        source.write_text(json.dumps([question_to_dict(question)]), encoding="utf-8")
        sources.append(source)
    application.import_library(str(sources[0]))
    first_bank = application.library.id
    application.scheduler.record("1", True)
    application.import_library(str(sources[1]))
    assert application.library.id != first_bank
    assert application.scheduler.states["1"].attempts == 0
    application.scheduler.record("1", False)
    application.import_library(str(sources[0]))
    assert application.scheduler.states["1"].correct == 1
    assert application.scheduler.states["1"].incorrect == 0
    assert application.library_store.load().id == first_bank


def test_import_report_lists_errors_without_replacing_active_bank(application, tmp_path):
    source = tmp_path / "invalid.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "题目", "A", "B", "C", "正确选项"])
    workbook.active.append([1, None, "a", "b", None, "D"])
    workbook.active.append([1, "duplicate", "a", "b", "c", "A"])
    workbook.save(source)
    before = application.library
    application.import_library(str(source))
    assert application.library == before
    assert application.scheduler.bank_id == before.id
    report = application.import_error_dialog
    assert report.isVisible()
    assert "4 处" in report.caption.text()
    assert "第 2 行" in report.text.toPlainText()
    assert "第 3 行" in report.text.toPlainText()


def test_format_and_extension_errors_are_reported_together(tmp_path):
    from app.models import ImportValidationError

    source = tmp_path / "mixed-errors.xlsx"
    workbook = Workbook()
    workbook.active.append(["ID", "题目", "A", "B", "C", "正确选项", "扩展"])
    workbook.active.append([1, "bad answer", "a", "b", "c", "F", None])
    workbook.active.append([2, "bad link", "a", "b", "c", "A", "url:javascript:bad"])
    workbook.save(source)
    with pytest.raises(ImportValidationError) as caught:
        LibraryStore(tmp_path / "cache/library.json").prepare_import(source)
    assert len(caught.value.issues) == 2
    assert {issue.location for issue in caught.value.issues} == {"第 2 行", "第 3 行"}


def test_saved_template_can_be_filled_and_imported(application, tmp_path, monkeypatch):
    output = tmp_path / "my-bank.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *args: (str(output), ""))
    application.settings.save_template()
    assert output.read_bytes() == (ROOT / "examples/question-template.xlsx").read_bytes()
    workbook = load_workbook(output)
    blank, guide = workbook["轮盘题库"], workbook["填写说明"]
    for index, row in enumerate(guide.iter_rows(min_row=10, max_row=12, values_only=True), 2):
        for column, value in enumerate(row, 1):
            blank.cell(index, column, value)
    workbook.save(output)
    workbook.close()
    imported = LibraryStore(tmp_path / "cache/library.json").import_file(output)
    assert len(imported.questions) == 3
    assert imported.questions[1].mode == "multiple"
    assert imported.questions[0].subtitle == "/ˈæpl/ n."
