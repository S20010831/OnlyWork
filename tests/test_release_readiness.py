import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.models import Option, QuestionFormatError, QuizQuestion
from app.paths import application_directory, resource_path
from app.questions import load_questions_from_xlsx

ROOT = Path(__file__).resolve().parents[1]


def test_domain_services_import_without_qt():
    code = """
import sys
class RejectQt:
    def find_spec(self, fullname, path=None, target=None):
        if fullname == "PySide6" or fullname.startswith("PySide6."):
            raise ImportError("Domain services must not import Qt")
sys.meta_path.insert(0, RejectQt())
from app.services import AppServices
from app.library import LibraryStore
from app.session import QuizSession
AppServices()
"""
    result = subprocess.run([sys.executable, "-c", code], cwd=ROOT, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_frozen_paths_keep_data_beside_executable_and_resources_in_bundle(tmp_path, monkeypatch):
    executable = tmp_path / "OnlyWork" / "OnlyWork.exe"
    bundle = executable.parent / "_internal"
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(executable))
    monkeypatch.setattr(sys, "_MEIPASS", str(bundle), raising=False)
    assert application_directory() == executable.parent
    assert resource_path("icon/onlywork.png") == bundle / "icon" / "onlywork.png"


def test_question_copies_mutable_inputs():
    options = [Option("A", "a"), Option("B", "b"), Option("C", "c")]
    correct = {"A"}
    question = QuizQuestion("Question", "", options, correct)
    original_id = question.id
    options.clear()
    correct.clear()
    assert len(question.options) == 3
    assert question.correct_ids == frozenset({"A"})
    assert question.id == original_id


def test_question_rejects_malformed_option_objects():
    with pytest.raises(QuestionFormatError, match="Option"):
        QuizQuestion("Question", "", ({"id": "A"}, Option("B", "b"), Option("C", "c")), {"A"})


def test_excel_finds_question_sheet_when_vocabulary_sheet_is_selected(tmp_path):
    workbook = Workbook()
    quiz = workbook.active
    quiz.title = "轮盘题库"
    quiz.append(["题目", "A", "B", "C", "正确选项"])
    quiz.append(["word", "a", "b", "c", "B"])
    vocabulary = workbook.create_sheet("词汇表")
    vocabulary.append(["单词", "释义"])
    workbook.active = 1
    path = tmp_path / "bank.xlsx"
    workbook.save(path)
    workbook.close()
    assert load_questions_from_xlsx(path)[0].prompt == "word"


def test_excel_rejects_ambiguous_question_sheets(tmp_path):
    workbook = Workbook()
    for name in ("Quiz A", "Quiz B"):
        sheet = workbook.create_sheet(name)
        sheet.append(["题目", "A", "B", "C", "正确选项"])
        sheet.append([name, "a", "b", "c", "B"])
    path = tmp_path / "ambiguous.xlsx"
    workbook.save(path)
    workbook.close()
    with pytest.raises(QuestionFormatError, match="多个题库工作表"):
        load_questions_from_xlsx(path)


def test_bundled_cet4_example_is_complete_and_importable(tmp_path):
    from app.library import LibraryStore

    store = LibraryStore(tmp_path / "library.json")
    imported = store.import_file(ROOT / "examples" / "cet4-vocabulary.xlsx")
    assert len(imported.questions) == 3848
    assert len({question.id for question in imported.questions}) == 3848
    assert [question.id for question in imported.questions] == [str(i) for i in range(1, 3849)]
    assert imported.questions[0].prompt == "abandon"
    assert imported.questions[0].subtitle == "/ə'bændən/ vt."
    assert all(
        len(question.options) == 4 and question.extension.type == "info"
        for question in imported.questions
    )
    assert store.load() == imported
    assert json.loads(store.path.read_text(encoding="utf-8"))["version"] == 1
