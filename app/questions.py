"""Built-in importers. Additional file formats register through QuestionImporters."""

import json
from pathlib import Path
from typing import Callable
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import (
    ExtensionAction,
    ImportIssue,
    ImportValidationError,
    Option,
    QuestionFormatError,
    QuizQuestion,
    validate_options,
)

EXCEL_REQUIRED_COLUMNS = {"题目", "a", "b", "c", "正确选项"}
EXCEL_COLUMNS = EXCEL_REQUIRED_COLUMNS | {"d", "e", "f", "id", "扩展"}
JSON_FIELDS = {"id", "prompt", "options", "correct_ids", "mode", "extension"}

# Re-export the models for existing imports of app.questions.
__all__ = [
    "Option",
    "QuizQuestion",
    "ExtensionAction",
    "QuestionFormatError",
    "load_questions_from_xlsx",
    "load_questions_from_json",
    "QuestionImporters",
    "default_importers",
    "SAMPLE_QUESTIONS",
]


def parse_serial_id(value: object) -> str:
    if value is None or str(value).strip() == "":
        return ""
    if isinstance(value, bool):
        raise QuestionFormatError("ID 必须为正整数序号")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    if not text.isascii() or not text.isdigit() or int(text) < 1:
        raise QuestionFormatError("ID 必须为正整数序号")
    return str(int(text))


def parse_correct(value: object, options: tuple[Option, ...], row: int = 0) -> frozenset[str]:
    if not isinstance(value, str):
        raise QuestionFormatError(f"第 {row} 行：正确选项为空或不是文本")
    result = frozenset(item.strip().upper() for item in value.split("|") if item.strip())
    if not result or not result <= {option.id for option in options}:
        raise QuestionFormatError(f"第 {row} 行：正确选项不存在或为空")
    return result


def parse_extension(value: object, row: int = 0) -> ExtensionAction | None:
    if value is None or str(value).strip() == "":
        return None
    raw = str(value).strip()
    if ":" not in raw:
        raise QuestionFormatError(f"第 {row} 行：扩展格式应为 type:value")
    name, content = raw.split(":", 1)
    if not name.strip() or not content.strip():
        raise QuestionFormatError(f"第 {row} 行：扩展类型和内容不能为空")
    return ExtensionAction(name.strip().lower(), content.strip())


def parse_question_text(value: object) -> tuple[str, str]:
    """The first line is the title; remaining lines are smaller supporting text."""
    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    prompt, _, subtitle = text.partition("\n")
    return prompt.strip(), subtitle.strip()


def question_to_dict(question: QuizQuestion) -> dict:
    return {
        "id": question.id,
        "prompt": "\n".join(part for part in (question.prompt, question.subtitle) if part),
        "options": [{"id": option.id, "text": option.text} for option in question.options],
        "correct_ids": sorted(question.correct_ids),
        "mode": question.mode,
        "extension": (
            {"type": question.extension.type, "value": question.extension.value}
            if question.extension
            else None
        ),
    }


def question_from_dict(raw: object) -> QuizQuestion:
    if not isinstance(raw, dict):
        raise QuestionFormatError("题目必须是对象")
    unknown = raw.keys() - JSON_FIELDS
    if unknown:
        raise QuestionFormatError("不支持的题目字段：" + "、".join(sorted(unknown)))
    try:
        if not isinstance(raw["options"], list) or not all(
            isinstance(option, dict) for option in raw["options"]
        ):
            raise QuestionFormatError("options 必须是选项对象数组")
        if not isinstance(raw["correct_ids"], list) or not all(
            isinstance(value, str) for value in raw["correct_ids"]
        ):
            raise QuestionFormatError("correct_ids 必须是编号数组")
        options = tuple(Option(option["id"], option["text"]) for option in raw["options"])
        correct = frozenset(raw["correct_ids"])
        extension = raw.get("extension")
        if extension is not None and not isinstance(extension, dict):
            raise QuestionFormatError("扩展必须是对象或 null")
        action = ExtensionAction(extension["type"], extension["value"]) if extension else None
        if not isinstance(raw["prompt"], str):
            raise QuestionFormatError("题目必须是文本")
        prompt, subtitle = parse_question_text(raw["prompt"])
        return QuizQuestion(
            prompt,
            subtitle,
            options,
            correct,
            raw.get("mode", "multiple" if len(correct) > 1 else "single"),
            action,
            raw.get("id", ""),
        )
    except (KeyError, TypeError, AttributeError) as error:
        raise QuestionFormatError("题目字段缺失或类型不正确") from error


def validate_unique(questions: tuple[QuizQuestion, ...]) -> tuple[QuizQuestion, ...]:
    if not questions:
        raise QuestionFormatError("题库没有题目")
    if len({question.id for question in questions}) != len(questions):
        raise QuestionFormatError("题目 ID 重复；相同题目请填写不同的 ID")
    return questions


def question_sheet(workbook):
    """Use the selected quiz sheet, or the only sheet with a quiz header."""

    def contains_questions(sheet):
        header = next(sheet.iter_rows(max_row=1, values_only=True), ())
        names = {str(value).strip().lower() for value in header if value is not None}
        return EXCEL_REQUIRED_COLUMNS <= names

    if contains_questions(workbook.active):
        return workbook.active
    candidates = [sheet for sheet in workbook if contains_questions(sheet)]
    if len(candidates) > 1:
        raise QuestionFormatError("存在多个题库工作表，请在 Excel 中选中要导入的工作表后保存")
    return candidates[0] if candidates else workbook.active


def load_questions_from_xlsx(path: str | Path) -> tuple[QuizQuestion, ...]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, ValueError, KeyError) as error:
        raise QuestionFormatError("无法读取 Excel，请检查文件是否为有效的 .xlsx 题库") from error
    try:
        rows = question_sheet(workbook).iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise QuestionFormatError("题库缺少表头")
        names = [
            str(value).strip().lower()
            for value in header
            if value is not None and str(value).strip()
        ]
        if len(set(names)) != len(names):
            raise QuestionFormatError("题库表头存在重复列")
        columns = {
            str(value).strip().lower(): index
            for index, value in enumerate(header)
            if value is not None
        }
        missing = EXCEL_REQUIRED_COLUMNS - columns.keys()
        if missing:
            raise QuestionFormatError("缺少必需列：" + "、".join(sorted(missing)))
        unsupported = columns.keys() - EXCEL_COLUMNS
        if unsupported:
            raise QuestionFormatError("不支持的题库列：" + "、".join(sorted(unsupported)))
        questions = []
        issues = []
        identifiers = {}
        for number, row in enumerate(rows, start=2):
            if all(value is None or str(value).strip() == "" for value in row):
                continue

            def cell(name: str) -> object:
                index = columns.get(name)
                return row[index] if index is not None and index < len(row) else None

            row_issues = []

            def check(parser, fallback):
                try:
                    return parser()
                except QuestionFormatError as error:
                    message = str(error).removeprefix(f"第 {number} 行：")
                    row_issues.append(ImportIssue(f"第 {number} 行", message))
                    return fallback

            prompt, subtitle = parse_question_text(cell("题目"))
            if not prompt:
                row_issues.append(ImportIssue(f"第 {number} 行", "题目不能为空"))
            options = tuple(
                Option(letter, str(cell(letter.lower())).strip())
                for letter in "ABCDEF"
                if cell(letter.lower()) is not None and str(cell(letter.lower())).strip()
            )
            check(lambda: validate_options(options), None)
            correct = check(lambda: parse_correct(cell("正确选项"), options, number), frozenset())
            action = check(lambda: parse_extension(cell("扩展"), number), None)
            identifier = check(lambda: parse_serial_id(cell("id")), "")
            if identifier:
                if identifier in identifiers:
                    row_issues.append(
                        ImportIssue(
                            f"第 {number} 行",
                            f"ID {identifier} 重复，首次出现在第 {identifiers[identifier]} 行",
                        )
                    )
                else:
                    identifiers[identifier] = number
            if row_issues:
                issues.extend(row_issues)
                continue
            mode = "multiple" if len(correct) > 1 else "single"
            question = QuizQuestion(
                prompt, subtitle, options, correct, mode, action, identifier, source_row=number
            )
            if not identifier and question.id in identifiers:
                issues.append(ImportIssue(f"第 {number} 行", "题目内容重复，请填写不同的 ID"))
            else:
                identifiers[question.id] = number
                questions.append(question)
        if issues:
            raise ImportValidationError(tuple(issues), tuple(questions))
        return validate_unique(tuple(questions))
    except (QuestionFormatError, OSError):
        raise
    except Exception as error:
        raise QuestionFormatError(f"Excel 内容损坏或格式无效：{error}") from error
    finally:
        workbook.close()


def load_questions_from_json(path: str | Path) -> tuple[QuizQuestion, ...]:
    try:
        raw = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (ValueError, UnicodeError) as error:
        raise QuestionFormatError("JSON 文件格式无效") from error
    if isinstance(raw, dict):
        if raw.get("version", 1) != 1:
            raise QuestionFormatError("不支持的题库版本")
        raw = raw.get("questions")
    if not isinstance(raw, list):
        raise QuestionFormatError("JSON 题库必须是题目数组或包含 questions 数组的对象")
    questions = []
    issues = []
    identifiers = {}
    for number, value in enumerate(raw, start=1):
        try:
            question = question_from_dict(value)
            if question.id in identifiers:
                raise QuestionFormatError(f"ID 重复，首次出现在第 {identifiers[question.id]} 道题")
            identifiers[question.id] = number
            questions.append(question)
        except QuestionFormatError as error:
            issues.append(ImportIssue(f"第 {number} 道题", str(error)))
    if issues:
        raise ImportValidationError(tuple(issues), tuple(questions))
    return validate_unique(tuple(questions))


QuestionLoader = Callable[[Path], tuple[QuizQuestion, ...]]


class QuestionImporters:
    def __init__(self) -> None:
        self._loaders: dict[str, QuestionLoader] = {}

    @property
    def file_filter(self) -> str:
        patterns = " ".join("*" + suffix for suffix in self._loaders)
        return f"题库文件 ({patterns})"

    def register(self, suffix: str, loader: QuestionLoader) -> None:
        suffix = "." + suffix.lower().lstrip(".")
        if suffix in self._loaders:
            raise ValueError(f"导入器已注册：{suffix}")
        self._loaders[suffix] = loader

    def load(self, path: Path) -> tuple[QuizQuestion, ...]:
        try:
            loader = self._loaders[path.suffix.lower()]
        except KeyError as error:
            raise QuestionFormatError(f"不支持的题库文件类型：{path.suffix}") from error
        try:
            return validate_unique(tuple(loader(path)))
        except (OSError, QuestionFormatError):
            raise
        except Exception as error:
            raise QuestionFormatError(f"题库读取失败：{error}") from error


def default_importers() -> QuestionImporters:
    importers = QuestionImporters()
    importers.register(".xlsx", load_questions_from_xlsx)
    importers.register(".json", load_questions_from_json)
    return importers


SAMPLE_QUESTIONS = (
    QuizQuestion(
        "abandon",
        "/əˈbændən/  v.",
        (Option("A", "放弃；遗弃"), Option("B", "记住；回忆"), Option("C", "逃跑；逃脱")),
        frozenset({"A"}),
    ),
    QuizQuestion(
        "Which are mammals?",
        "可多选",
        (
            Option("A", "鲸鱼"),
            Option("B", "蝙蝠"),
            Option("C", "鳄鱼"),
            Option("D", "海豚"),
            Option("E", "企鹅"),
        ),
        frozenset({"A", "B", "D"}),
        "multiple",
        ExtensionAction("info", "鲸鱼、蝙蝠和海豚是哺乳动物。"),
    ),
)
